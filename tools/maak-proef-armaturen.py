"""Maakt tools/proef-armaturen.xlsx: de proeflijst waarmee tools/controleer-logica.mjs
het inlezen van een Excel-lijst in het armaturenboek en de railtool controleert
(armaturen-excel.js). Gemaakt met openpyxl, zodat het een echt .xlsx is - tekst in de
cel, een getal als artikelcode, opgemaakte tekst in één cel - en geen nagebootst
bestand.

De opmaak is die van de lijst die inside sales zelf maakt (oktober 2026): drie
kopregels (Armatuurcode / Aantallen per ruimte en verdieping / TOTAAL), kolom A de
code, F het totale aantal, I de artikelcode - accessoires erachter met een + - J de
omschrijving op dezelfde manier, L de opmerking. B tot en met E, G, H en K zijn ruis
die genegeerd moet worden, en onderaan staan regels met alleen een 0 in de
totaalkolom, zoals in de echte lijst. Geen prijzen: de repo is openbaar.

    python3 tools/maak-proef-armaturen.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

uit = Path(__file__).with_name('proef-armaturen.xlsx')

MONDIAL_120 = ('Pragmalux LED Downlight Mondial IP54 facet reflector max. 16W max. 1800lm 3000K-3500K-4000K 3-CCT '
               'CRI>90 UGR<19 60D Ø120 excl. LED Driver en frame')
DURA_SN = ('Pragmalux LED Plafonnière / Wandarmatuur Dura-M Wit Ø340 IP64 IK10 8-18W 750-2075lm 2200K-2700K 2-CCT '
           '+ bewegingssensor + noodmodule 1uur autotest')
DURA_S = ('Pragmalux LED Plafonnière / Wandarmatuur Dura-M Wit Ø340 IP64 IK10 8-18W 750-2075lm 2200K-2700K 2-CCT '
          '+ bewegingssensor')
SIGMA = ('Pragmalux LED Paneel 30x120cm Sigma G2 IP44 18-26W 3000K-4000K 2-CCT 2400-3650lm UGR<19 '
         '+Osram Driver (2x28W)')
ESSENCE = ('Pragmalux LED Downlight Essence G2 IP54 wit standaard max. 15W 3000K-6000K 3-CCT max. 1650lm 100D '
           'Ø200 Buitenmaat - Gatmaat Ø150 excl. LED Driver')
LUNA = ('Pragmalux LED Inbouw/Opbouw Downlight Luna G2 IP44 12W/18W 3000K-6000K 3-CCT 1400-2050lm Ø217 '
        'Buitenmaat - Gatmaat Ø65-185 incl. LED Driver')
NOOD = 'Pragmalux Noodmodule Kit t.b.v. Panelen & Downlights 3 uur met autotest'
DRIVER = 'Philips Driver Xitanium 20W/m 300mA 44V DS I 230V incl. DC kabel'

wb = Workbook()
ws = wb.active
ws.title = 'Table 1'

# de kop zoals in de echte lijst, met de tikfout in OMSCRHIJVING
for k, v in {'A1': 'Armatuurcode', 'B1': 'Aantallen', 'G1': 'Leverancier / merk / type ', 'H1': 'Voorbeeld',
             'I1': 'ARTIKELCODE', 'J1': 'OMSCRHIJVING', 'K1': 'VOORBEELD', 'L1': 'OPMERKINGEN',
             'B2': 'KANTOOR', 'E2': 'HAL', 'B3': 'BG', 'C3': 'V1', 'D3': 'V2', 'F3': 'TOTAAL'}.items():
    ws[k] = v
ws.merge_cells('B1:F1')
ws.merge_cells('B2:D2')

rijen = {
    # gelijk aan de catalogus; de artikelcode is een getal
    4: {'A': 'A', 'B': 12, 'F': 12, 'G': 'Bestekarmatuur X', 'I': 1092379, 'J': MONDIAL_120},
    # één woord anders dan de catalogus (3uur, de catalogus zegt 1uur), met een opmerking
    5: {'A': 'B', 'B': 4, 'F': 4, 'I': '1084459-SN', 'J': DURA_SN.replace('1uur', '3uur'), 'L': 'Niet IK10'},
    # niet in de catalogus; een & in de tekst
    6: {'A': 'C', 'E': 2, 'F': 2, 'I': '9999999', 'J': 'Eigen armatuur & zo', 'L': 'via de leverancier'},
    # geen omschrijving in de lijst, geen aantal
    7: {'A': 'D', 'I': '1047515'},
    # een tussenkop: alleen kolom A en een 0 in de totaalkolom
    8: {'A': 'Gang', 'F': 0},
    # niet in de catalogus en geen omschrijving
    9: {'A': 'E', 'I': '8888888'},
    # gelijk, met spaties eromheen en een dubbele spatie
    11: {'A': ' G ', 'I': ' 1092416 ',
         'J': 'Pragmalux LED Downlight Mondial IP54  platte frame alu wit Ø168 buitenmaat - gatmaat Ø150'},
    # een armatuur met een accessoire; de naam van het armatuur heeft zelf een + ("+Osram")
    12: {'A': 'H', 'B': 3, 'C': 5, 'F': 8, 'I': '1000947-OS + 1000282', 'J': SIGMA + ' + ' + NOOD},
    # twee accessoires
    13: {'A': 'K', 'D': 4, 'F': 4, 'I': 'WS1046761 + 9322219-300 + 1000282', 'J': ESSENCE + ' + ' + DRIVER + ' + ' + NOOD},
    # een armatuur met " + " in zijn eigen naam, en een accessoire dat de catalogus kent
    14: {'A': 'M', 'F': 6, 'I': '1084459-S+1047515', 'J': DURA_S + ' + ' + LUNA},
    # een special
    15: {'A': 'S', 'F': 3, 'I': 'Special', 'J': 'DLC - volledige eigen naam special'},
}
for r, cellen in rijen.items():
    for k, v in cellen.items():
        ws[f'{k}{r}'] = v
# gelijk, maar de omschrijving is opgemaakte tekst: drie stukken in één cel
ws['A10'] = 'F'
ws['I10'] = '1092409'
ws['J10'] = CellRichText('Pragmalux LED Downlight Mondial IP54 ', TextBlock(InlineFont(b=True), 'platte frame'),
                         ' alu wit Ø132 buitenmaat - gatmaat Ø120')
# lege regels van het sjabloon: alleen een 0 in de totaalkolom
for r in range(16, 21):
    ws[f'F{r}'] = 0

# een tweede blad dat niet gelezen mag worden
ws2 = wb.create_sheet('Notities')
ws2['A1'] = 'X'
ws2['I1'] = '1092423'
ws2['J1'] = 'hoort er niet bij'

wb.save(uit)
print(uit)
